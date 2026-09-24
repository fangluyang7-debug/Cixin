from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DELIVERABLES = ROOT / "docs" / "deliverables"
BACKUP_DIR = ROOT / ".tmp" / "docx-backups"
OUT_PATH = next(
    path
    for path in DELIVERABLES.glob("*.docx")
    if path.name.startswith("AI") and not path.name.startswith("~$")
)


def set_east_asia_font(run, font_name: str, size: float | None = None, bold: bool | None = None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_style_font(style, font_name: str, size: float, bold: bool = False):
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    style._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)


def paragraph(text: str = "", style: str | None = None, align=None, before=0, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        if style in {"Heading 1", "Heading 2"}:
            set_east_asia_font(r, "黑体", 14 if style == "Heading 1" else 12, True)
        else:
            set_east_asia_font(r, "宋体", 11, False)
    return p


def heading(text: str, level: int = 1):
    style = "Heading 1" if level == 1 else "Heading 2"
    p = paragraph(text, style=style, before=10 if level == 1 else 6, after=5)
    return p


def body(text: str):
    return paragraph(text, style="Normal", after=5)


def bullet(text: str):
    p = paragraph(style="Normal", after=3)
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    marker = p.add_run("· ")
    set_east_asia_font(marker, "宋体", 11, False)
    r = p.add_run(text)
    set_east_asia_font(r, "宋体", 11, False)
    return p


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 10.5, font: str = "宋体"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(text)
    set_east_asia_font(r, font, size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "BFBFBF")


def table(headers: list[str], rows: list[list[str]], widths_cm: list[float]):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t)
    for idx, header in enumerate(headers):
        cell = t.rows[0].cells[idx]
        cell.width = Cm(widths_cm[idx])
        shade_cell(cell, "EDEDED")
        set_cell_text(cell, header, bold=True, size=10.5, font="黑体")
    for row in rows:
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].width = Cm(widths_cm[idx])
            set_cell_text(cells[idx], value, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return t


BACKUP_DIR.mkdir(parents=True, exist_ok=True)
if OUT_PATH.exists():
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy2(OUT_PATH, BACKUP_DIR / f"AI使用总结文档-garbled-backup-{stamp}.docx")

doc = Document()
section = doc.sections[0]
section.start_type = WD_SECTION_START.NEW_PAGE
section.top_margin = Cm(2.35)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.45)
section.right_margin = Cm(2.45)

set_style_font(doc.styles["Normal"], "宋体", 11, False)
set_style_font(doc.styles["Heading 1"], "黑体", 14, True)
set_style_font(doc.styles["Heading 2"], "黑体", 12, True)

title = paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
r = title.add_run("AI 拍照识物与智能比价购物助手")
set_east_asia_font(r, "黑体", 18, True)
subtitle = paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
r = subtitle.add_run("AI 使用总结文档")
set_east_asia_font(r, "黑体", 16, True)
date = paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
r = date.add_run("2026年6月10日")
set_east_asia_font(r, "宋体", 11, False)

body(
    "本文档说明本项目在实际完成过程中如何使用 AI 工具辅助开发、调试和交付。"
    "文档重点不重复架构设计文档的模块说明，而是记录 Codex、Claude Code、Antigravity IDE 等工具"
    "在前端、后端、数据处理、Prompt 设计、部署测试和文档整理中的具体作用，以及人工确认、成本控制和安全边界。"
)
body(
    "项目的核心目标是构建一个可拍照识物、跨平台召回商品、支持自然语言继续筛选的智能比价购物助手。"
    "由于涉及真实商品数据、图片 embedding、模型调用、移动端打包和云端服务部署，AI 工具主要承担"
    "工程协作、代码阅读、方案生成、脚本编写和问题定位工作，最终功能仍通过人工审查、运行测试和业务校验确定。"
)

heading("1 文档定位与使用原则")
body(
    "AI 在本项目中不是单一的聊天问答工具，而是贯穿项目推进的工程助手。开发过程中将 AI 用于快速阅读代码、"
    "生成候选实现、补充测试脚本、分析日志、整理接口说明和优化文档表达。对涉及数据安全、云端成本、真实商品导入"
    "和环境变量的任务，均采用人工确认后再执行的方式，避免把敏感信息直接暴露给模型或让模型自动决定高成本操作。"
)
bullet("以人工目标为主：先明确业务目标和演示要求，再让 AI 辅助拆解任务、定位文件和给出实现建议。")
bullet("以项目代码为准：AI 生成的结论必须回到实际仓库文件、运行结果、接口响应和数据库状态中验证。")
bullet("以可替换为设计目标：模型、搜索、存储、商品数据源和 Adapt 层均尽量通过接口隔离，减少后续替换成本。")
bullet("以真实数据为核心价值：数据来自聚合 API 和实际采集文件，受成本限制仅生成部分双 embedding 演示数据，但数据真实性更接近实用场景。")

heading("2 AI 工具分工")
table(
    ["工具", "主要使用场景", "在项目中的具体产出"],
    [
        [
            "Codex",
            "深度阅读仓库、修改代码、运行命令、生成脚本、检查构建结果。",
            "协助梳理 NestJS 后端、Flutter App、商品池导入脚本、Zeabur 环境变量、APK 打包流程；定位上传失败、embedding 数量异常和拍照搜索慢等问题。",
        ],
        [
            "Claude Code",
            "辅助进行局部代码推理、Prompt 文案打磨、复杂逻辑解释和可读性优化。",
            "用于讨论自然语言筛选、多轮对话状态、Prompt 输出格式和文档表达方式，帮助把工程实现解释成答辩可理解的内容。",
        ],
        [
            "Antigravity IDE",
            "在开发环境中辅助查看文件、组织上下文、对比修改和进行多文件理解。",
            "用于配合前后端联调、查看页面状态、理解项目目录，以及在开发过程中快速切换前端、后端和文档材料。",
        ],
        [
            "通用大模型能力",
            "生成结构化 Prompt、检查文档逻辑、提出可插拔设计和性能优化建议。",
            "形成模型任务编排、Adapter 层解耦、商品池构建和安全说明等内容，并对答辩材料进行语言组织。",
        ],
    ],
    [2.7, 5.1, 7.3],
)

heading("3 前端开发中的 AI 辅助")
body(
    "前端包括 Flutter 移动端 App 和用于评委/调试展示的 Web 页面。AI 首先帮助识别真正需要打包的是"
    " apps/mobile-flutter 中的 Flutter App，而不是早期用于调试的 judge-entry-web。随后在移动端流程中，"
    "辅助梳理了拍照、相册选择、主体框选择、图片压缩、上传图片、创建会话、候选商品展示和继续筛选等链路。"
)
bullet("在 App 打包阶段，AI 帮助区分 Capacitor 调试网站和 Flutter 正式 App，给出 release APK 的构建参数。")
bullet("在拍照搜索阶段，AI 辅助分析移动端图片压缩参数，降低上传图片尺寸，减少拍照后的等待时间。")
bullet("在前端展示阶段，AI 帮助检查候选商品卡片、商品 URL、店铺、平台、价格和图片字段是否与后端接口适配。")
bullet("在调试体验阶段，AI 曾帮助搭建 Adapt 可视化页面，用于批量查看裁剪效果；后续根据项目目标回退为正式入库链路优化。")

heading("4 后端开发中的 AI 辅助")
body(
    "后端采用 NestJS + Prisma 组织，包含图片资产、查询会话、候选商品、商品池、embedding、记忆系统和维护接口等模块。"
    "AI 在后端开发中主要承担代码阅读、链路追踪和低风险改造任务。例如，当拍照搜索变慢时，AI 通过阅读会话创建、"
    "QueryImagePreprocess、EmbeddingProvider 和 ANN 检索服务，确认耗时不只是相似度计算，还包括图片上传、裁剪、"
    "embedding 生成、候选融合和数据库读取。"
)
bullet("在 Controller 与 Service 层，AI 协助梳理接口输入输出，明确 App 接口与维护接口的权限边界。")
bullet("在 Provider 与 Adapter 层，AI 协助识别模型、搜索、存储、商品导入、图片预处理等可替换点。")
bullet("在 Prisma 层，AI 协助补充复合索引，优化商品池查询、embedding 读取和候选快照访问。")
bullet("在异常处理上，AI 协助分析 Bad Gateway、超时、批次失败、平台字段异常和 embedding 数量不一致等问题。")

heading("5 数据处理与商品池建设中的 AI 辅助")
body(
    "商品数据来自聚合 API 和多平台采集文件，包含 JSON 与 Excel 两类格式。AI 辅助完成了数据字段理解、格式归一、"
    "分组切分、去重检查、平台识别和上传脚本整理。项目强调真实商品数据的价值：真实数据包含商品图片、价格、平台、"
    "店铺名称、跳转 URL 和平台差异，能更真实地验证拍照识物与比价流程，而不是只在模拟数据上演示。"
)
body(
    "由于双 embedding 会显著增加模型调用费用和入库时间，项目没有对所有商品强制生成 visual + multimodal 双向量。"
    "实际策略是：保留一部分双 embedding 商品用于演示模型能力和方案可行性，后续大规模上传以 visual embedding 为主，"
    "在成本可控的前提下保证检索链路能正常运行。"
)
table(
    ["数据任务", "AI 辅助内容", "人工确认重点"],
    [
        ["JSON/Excel 归一", "分析字段含义，转换为后端标准商品结构。", "确认平台、URL、图片、价格和店铺字段没有错位。"],
        ["分组上传", "生成 1000 组、50 条 batch 的切分与上传脚本。", "避免重复上传，检查已上传批次和失败批次。"],
        ["平台识别", "定位苏宁、唯品会、闲鱼等被误写为 manual 的风险。", "先迁移或删除异常批次，再导入修正版数据。"],
        ["成本控制", "分析 embedding、COS、模型调用的资源消耗。", "决定只生成必要的 embedding 类型，控制演示规模。"],
    ],
    [3.0, 7.0, 5.1],
)

heading("6 Prompt 设计与模型任务编排")
body(
    "项目中的 AI 调用并不是简单把用户一句话直接交给模型，而是围绕多轮购物决策设计了结构化 Prompt 和任务边界。"
    "例如用户继续输入“便宜一点”“只看苏宁”“不要这个品牌”时，模型需要输出 intent、filterPatch、filterRemove、"
    "assistantMessage、confidence 和 raw 等结构化字段，后端再根据这些字段决定是改筛选条件、重置条件还是继续解释商品。"
)
bullet("Prompt 输出采用 JSON 结构，便于后端校验和容错，减少自然语言结果不可控的问题。")
bullet("将商品识别、类别判断、筛选理解、候选解释和结果排序拆成不同任务，避免一个 Prompt 同时承担过多职责。")
bullet("对模型结果加入 Guard 与 fallback：当解析失败、置信度不足或模型超时时，系统可以退回已有筛选条件或基础候选结果。")
bullet("Prompt 设计强调鲁棒性：即使商品信息不完整，也尽量输出可解释、可继续追问的结构化结果。")

heading("7 可插拔设计与 AI 协作价值")
body(
    "本项目的一个核心评分点是可插拔设计。AI 在阅读代码和提出改造方案时，重点围绕“替换某个能力不影响主流程”展开。"
    "例如模型服务通过 ModelAdapter 隔离，embedding 通过 EmbeddingProvider 隔离，商品检索通过 SearchProvider 隔离，"
    "图片存储通过 StorageAdapter 隔离，商品导入通过 ProductImportAdapter 与 ProductImageEmbeddingAdapter 隔离。"
)
body(
    "这种设计使项目可以根据不同业务场景快速替换能力：演示时使用成本较低的 visual embedding；正式部署时可以接入更强的多模态模型；"
    "检索性能不足时可以从当前数据库扫描升级为 pgvector、Qdrant、Milvus 或其他向量数据库；对象存储也可以从本地 mock 切换为云端 COS。"
)
table(
    ["可替换部分", "当前实现", "可替换方向"],
    [
        ["模型识别", "OpenAI-compatible / 火山方舟等模型适配。", "替换为其他视觉模型、分类模型或本地模型。"],
        ["Embedding", "视觉 embedding 与部分多模态 embedding。", "替换为成本更低或召回更强的 embedding 模型。"],
        ["搜索链路", "商品池本地检索、标签融合、ANN 召回。", "替换为真实向量数据库、混合搜索或平台实时搜索。"],
        ["图片处理", "查询图裁剪、商品图标准化、Adapt 层。", "替换裁剪算法、检测模型或分辨率策略。"],
        ["数据源", "聚合 API 与采集文件导入。", "接入更多平台 API，或导入商家自有商品库。"],
    ],
    [3.0, 5.7, 6.4],
)

heading("8 调试、部署与性能优化中的 AI 辅助")
body(
    "在项目推进中，AI 参与了大量工程调试工作，包括 PowerShell 上传脚本、Zeabur 环境变量、维护 token、云端批次状态、"
    "商品池统计、embedding 数量查询、APK 打包和接口健康检查。AI 的价值在于可以把终端报错、后端日志和代码实现联系起来，"
    "快速判断是脚本参数、网络超时、模型调用、数据库字段还是服务器状态导致问题。"
)
bullet("当批量上传因超时中断时，AI 帮助修改脚本，使超时可继续，不因单个 batch 失败阻断整组任务。")
bullet("当拍照搜索变慢时，AI 帮助定位到 query embedding 生成、COS 往返和非索引化向量扫描等真实瓶颈。")
bullet("当 APK 拍照后无法连接时，AI 帮助区分本地调试地址与云端 API 地址，确认打包时需要写入 API_BASE_URL。")
bullet("当文档或报告格式不统一时，AI 帮助根据课程要求整理架构文档、API 文档和 AI 使用总结文档。")

heading("9 安全与环境变量管理")
body(
    "项目中包含维护 token、对象存储密钥、模型 API key 和数据库地址等敏感信息。AI 辅助开发时遵循最小暴露原则："
    "只解释环境变量的用途，不在文档中泄露具体值；命令示例通过读取本地 .env 或 Zeabur 环境变量传入 token，避免把密钥写死在脚本中。"
)
bullet("MAINTENANCE_API_TOKEN 用于保护商品池维护接口，避免未授权导入、删除或触发高成本模型任务。")
bullet("OBJECT_STORAGE_* 用于访问云端对象存储，负责保存原图、商品图和必要的资源文件。")
bullet("EMBEDDING_API_KEY、VISION_MODEL_API_KEY 等模型密钥只应保存在本地 .env 或 Zeabur 私有环境变量中。")
bullet("文档与代码提交中应避免出现真实密钥，示例命令只展示变量名和读取方式。")

heading("10 人工确认与 AI 边界")
body(
    "AI 提高了项目推进速度，但并不替代人工判断。涉及真实商品数据、成本、云端导入、模型选择和答辩材料时，仍需要开发者根据业务目标作最终决定。"
    "例如是否继续上传全量数据、是否生成双 embedding、是否保存裁剪图到 COS、是否接入真实向量数据库，这些选择都需要结合预算、演示要求和服务器性能确认。"
)
table(
    ["事项", "AI 可以完成", "必须人工确认"],
    [
        ["代码修改", "阅读上下文、生成补丁、运行构建检查。", "确认功能是否符合演示目标和用户体验。"],
        ["数据导入", "切分文件、生成上传命令、分析失败原因。", "确认数据来源合法、平台字段正确、不会重复导入。"],
        ["模型调用", "设计 Prompt、分析成本、优化参数。", "确认预算、密钥管理和线上调用频率。"],
        ["部署打包", "给出构建命令、检查接口地址和环境变量。", "确认云端服务已重启、APK 指向正确后端。"],
    ],
    [2.7, 6.2, 6.2],
)

heading("11 总结")
body(
    "本项目中 AI 工具的主要价值体现在工程协作和复杂链路整理：它帮助开发者快速理解前后端代码、生成脚本、定位问题、"
    "优化性能、补充可插拔设计说明，并把真实商品池、embedding 检索、多轮对话和移动端打包这些分散任务连接成可交付的系统。"
)
body(
    "从最终结果看，AI 的使用并不是为了替代项目本身的设计与实现，而是让开发者能更快发现问题、比较方案和完成验证。"
    "项目依然以真实数据、明确接口、可替换模型、可扩展商品池和安全环境变量管理作为核心交付价值。受成本和时间限制，"
    "真实数据规模与双 embedding 覆盖率没有无限扩大，但已经足以证明系统具备从真实商品数据出发构建智能比价助手的可行性。"
)

try:
    doc.save(OUT_PATH)
    print(OUT_PATH)
except PermissionError:
    fallback = OUT_PATH.with_name(f"{OUT_PATH.stem}_修复版{OUT_PATH.suffix}")
    doc.save(fallback)
    print(fallback)
